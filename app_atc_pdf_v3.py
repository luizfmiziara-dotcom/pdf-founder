from __future__ import annotations

import io
import re
import zipfile
from collections import Counter
from pathlib import Path

import fitz
import streamlit as st
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps


DIRECTION_PATTERN = re.compile(
    r"(?<!\d)(\d{1,3}(?:[.,]\d+)?)(?!\d)"
    r"\s*(?:º|°|GRAU|GRAUS)?\s*\(?\s*NM\s*\)?",
    flags=re.IGNORECASE,
)


def normalize_text(text: str) -> str:
    return " ".join(
        (text or "")
        .replace("\u00a0", " ")
        .replace("\n", " ")
        .split()
    )


def parse_angle(value: str) -> float:
    return float(value.replace(",", "."))


def format_number(angle: float) -> str:
    if float(angle).is_integer():
        return str(int(angle))

    return (
        f"{angle:.2f}"
        .rstrip("0")
        .rstrip(".")
        .replace(".", ",")
    )


def format_angle(angle: float) -> str:
    return f"{format_number(angle)}º (NM)"


def safe_stem(filename: str) -> str:
    stem = Path(filename).stem
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", stem).strip("_")
    return cleaned or "arquivo"


def extract_angles_from_text(text: str) -> list[float]:
    angles: list[float] = []

    for match in DIRECTION_PATTERN.finditer(normalize_text(text)):
        angle = parse_angle(match.group(1))
        if 0 <= angle <= 360:
            angles.append(angle)

    return sorted(set(angles))


def matching_pages(pdf_bytes: bytes) -> tuple[list[int], list[str]]:
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_found: list[int] = []
    evidence: list[str] = []

    try:
        for page_index, page in enumerate(document):
            angles = extract_angles_from_text(page.get_text("text"))

            if angles:
                pages_found.append(page_index)
                evidence.append(", ".join(format_angle(a) for a in angles))
    finally:
        document.close()

    return pages_found, evidence


def crop_photo_footer(image_bytes: bytes, crop_percent: float) -> bytes:
    with Image.open(io.BytesIO(image_bytes)) as source_image:
        image = ImageOps.exif_transpose(source_image)

        if image.mode != "RGB":
            image = image.convert("RGB")

        width, height = image.size
        pixels_to_remove = round(height * crop_percent / 100)
        new_height = height - pixels_to_remove

        if new_height <= 0:
            raise ValueError("Percentual de corte inválido para a fotografia.")

        cropped = image.crop((0, 0, width, new_height))
        output = io.BytesIO()
        cropped.save(output, format="JPEG", quality=95, optimize=True)
        return output.getvalue()


def is_probable_photo(block: dict, page_rect: fitz.Rect) -> bool:
    image_width = block.get("width", 0)
    image_height = block.get("height", 0)
    rect = fitz.Rect(block["bbox"])

    if image_width < 250 or image_height < 180:
        return False

    if rect.width < page_rect.width * 0.20:
        return False

    if rect.height < page_rect.height * 0.10:
        return False

    if rect.y0 < page_rect.height * 0.12:
        return False

    return True


def image_blocks(page: fitz.Page) -> list[dict]:
    page_dict = page.get_text("dict")

    blocks = [
        block
        for block in page_dict.get("blocks", [])
        if block.get("type") == 1
        and block.get("image")
        and is_probable_photo(block, page.rect)
    ]

    return sorted(
        blocks,
        key=lambda block: (
            round(fitz.Rect(block["bbox"]).y0, 1),
            fitz.Rect(block["bbox"]).x0,
        ),
    )


def angle_labels_with_positions(page: fitz.Page) -> list[dict]:
    labels: list[dict] = []
    text_dict = page.get_text("dict")

    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue

        for line in block.get("lines", []):
            line_text = "".join(
                span.get("text", "")
                for span in line.get("spans", [])
            )

            for match in DIRECTION_PATTERN.finditer(normalize_text(line_text)):
                angle = parse_angle(match.group(1))

                if not 0 <= angle <= 360:
                    continue

                bbox = fitz.Rect(line.get("bbox", block.get("bbox")))
                labels.append(
                    {
                        "angle": angle,
                        "rect": bbox,
                        "text": match.group(0),
                    }
                )

    return labels


def associate_labels_to_photos(page: fitz.Page) -> list[tuple[float, dict]]:
    photos = image_blocks(page)
    labels = angle_labels_with_positions(page)

    if not photos or not labels:
        return []

    remaining = list(range(len(photos)))
    associations: list[tuple[float, dict]] = []

    for label in sorted(labels, key=lambda item: (item["rect"].y0, item["rect"].x0)):
        if not remaining:
            break

        label_rect = label["rect"]
        label_x = (label_rect.x0 + label_rect.x1) / 2
        label_y = (label_rect.y0 + label_rect.y1) / 2

        def score(photo_index: int) -> tuple[float, float]:
            photo_rect = fitz.Rect(photos[photo_index]["bbox"])
            photo_x = (photo_rect.x0 + photo_rect.x1) / 2
            photo_y = (photo_rect.y0 + photo_rect.y1) / 2
            horizontal = abs(photo_x - label_x)
            vertical = photo_rect.y0 - label_rect.y1

            if vertical >= -5 and horizontal <= max(photo_rect.width, 120):
                return (0.0, vertical + horizontal * 0.20)

            distance = ((photo_x - label_x) ** 2 + (photo_y - label_y) ** 2) ** 0.5
            return (1.0, distance)

        best_index = min(remaining, key=score)
        remaining.remove(best_index)
        associations.append((label["angle"], photos[best_index]))

    return associations


def extract_photos(
    pdf_bytes: bytes,
    page_indexes: list[int],
    remove_timestamp: bool,
    crop_percent: float,
) -> list[dict]:
    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    extracted: list[dict] = []

    try:
        for page_index in page_indexes:
            page = document[page_index]
            associations = associate_labels_to_photos(page)

            if not associations:
                angles = extract_angles_from_text(page.get_text("text"))
                photos = image_blocks(page)
                associations = list(zip(angles, photos))

            for angle, block in associations:
                image_bytes = block.get("image")
                if not image_bytes:
                    continue

                if remove_timestamp:
                    image_bytes = crop_photo_footer(image_bytes, crop_percent)

                extracted.append(
                    {
                        "angle": float(angle),
                        "image": image_bytes,
                        "source_page": page_index + 1,
                    }
                )
    finally:
        document.close()

    extracted.sort(key=lambda item: (item["angle"], item["source_page"]))

    totals = Counter(item["angle"] for item in extracted)
    sequences: Counter = Counter()

    for item in extracted:
        angle = item["angle"]
        sequences[angle] += 1
        base_label = format_angle(angle)

        if totals[angle] > 1:
            item["label"] = f"{base_label} - Foto {sequences[angle]}"
        else:
            item["label"] = base_label

    return extracted


def prepare_excel_image(image_bytes: bytes) -> tuple[io.BytesIO, int, int]:
    with Image.open(io.BytesIO(image_bytes)) as source_image:
        image = ImageOps.exif_transpose(source_image)

        if image.mode != "RGB":
            image = image.convert("RGB")

        max_width = 430
        max_height = 290
        ratio = min(max_width / image.width, max_height / image.height)
        width = max(1, round(image.width * ratio))
        height = max(1, round(image.height * ratio))

        buffer = io.BytesIO()
        image.save(buffer, format="PNG", optimize=True)
        buffer.seek(0)
        return buffer, width, height


def create_photos_excel(filename: str, photos: list[dict]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "FOTOS 360°"
    worksheet.sheet_view.zoomScale = 50
    worksheet.sheet_view.showGridLines = False

    worksheet.column_dimensions["A"].width = 3
    worksheet.column_dimensions["B"].width = 58
    worksheet.column_dimensions["C"].width = 4
    worksheet.column_dimensions["D"].width = 58
    worksheet.column_dimensions["E"].width = 3

    worksheet.merge_cells("A1:E1")
    title = worksheet["A1"]
    title.value = "FOTOS 360°"
    title.font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    title.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    title.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 32

    worksheet.merge_cells("A3:E3")
    section = worksheet["A3"]
    section.value = "FOTOS PANORÂMICAS"
    section.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    section.fill = PatternFill(fill_type="solid", fgColor="5B9BD5")
    section.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[3].height = 24

    image_buffers: list[io.BytesIO] = []
    current_row = 5

    for photo_index in range(0, len(photos), 2):
        pair = photos[photo_index:photo_index + 2]

        for pair_index, photo in enumerate(pair):
            column = "B" if pair_index == 0 else "D"
            label_cell = worksheet[f"{column}{current_row}"]
            label_cell.value = photo.get("label", format_angle(photo["angle"]))
            label_cell.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
            label_cell.alignment = Alignment(horizontal="center", vertical="center")

            buffer, image_width, image_height = prepare_excel_image(photo["image"])
            image_buffers.append(buffer)

            excel_image = ExcelImage(buffer)
            excel_image.width = image_width
            excel_image.height = image_height
            worksheet.add_image(excel_image, f"{column}{current_row + 1}")

        worksheet.row_dimensions[current_row].height = 22

        for row in range(current_row + 1, current_row + 17):
            worksheet.row_dimensions[row].height = 16

        current_row += 18

    worksheet.merge_cells(
        start_row=current_row,
        start_column=1,
        end_row=current_row,
        end_column=5,
    )
    footer = worksheet.cell(row=current_row, column=1)
    footer.value = f"Total de fotografias: {len(photos)}"
    footer.font = Font(name="Calibri", size=10, italic=True, color="666666")
    footer.alignment = Alignment(horizontal="right", vertical="center")

    worksheet.freeze_panes = "A5"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_area = f"A1:E{current_row}"

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def process_file(
    filename: str,
    pdf_bytes: bytes,
    remove_timestamp: bool,
    crop_percent: float,
) -> dict:
    pages, evidence = matching_pages(pdf_bytes)

    if not pages:
        return {
            "filename": filename,
            "pages": [],
            "evidence": [],
            "photos": [],
            "excel_output": None,
        }

    photos = extract_photos(
        pdf_bytes=pdf_bytes,
        page_indexes=pages,
        remove_timestamp=remove_timestamp,
        crop_percent=crop_percent,
    )

    excel_output = create_photos_excel(filename, photos) if photos else None
    stem = safe_stem(filename)

    return {
        "filename": filename,
        "pages": pages,
        "evidence": evidence,
        "photos": photos,
        "excel_output": excel_output,
        "excel_name": f"{stem}_fotos_360.xlsx",
    }


st.set_page_config(
    page_title="Gerador de Excel de Fotos NM",
    page_icon="📷",
    layout="wide",
)

st.title("Gerador de Excel de fotos panorâmicas NM")
st.write(
    "Envie um ou mais PDFs. O aplicativo identifica qualquer ângulo entre "
    "**0 e 360 associado a NM**, extrai as fotografias e gera um Excel dinâmico "
    "com duas fotos por linha, ordenadas pelo ângulo."
)

st.subheader("Tratamento das fotografias")
remove_timestamp = st.checkbox(
    "Remover a faixa inferior com data, hora e coordenadas",
    value=False,
)

crop_percent = st.slider(
    "Percentual a remover da parte inferior de cada foto",
    min_value=1.0,
    max_value=20.0,
    value=8.0,
    step=0.5,
    disabled=not remove_timestamp,
)

uploaded_files = st.file_uploader(
    "Selecione os arquivos PDF",
    type=["pdf"],
    accept_multiple_files=True,
)

if uploaded_files:
    results: list[dict] = []
    progress = st.progress(0)
    status = st.empty()

    for index, uploaded in enumerate(uploaded_files, start=1):
        status.write(f"Processando {uploaded.name}...")

        try:
            result = process_file(
                filename=uploaded.name,
                pdf_bytes=uploaded.getvalue(),
                remove_timestamp=remove_timestamp,
                crop_percent=crop_percent,
            )
            results.append(result)
        except Exception as error:
            results.append({"filename": uploaded.name, "error": str(error)})

        progress.progress(index / len(uploaded_files))

    status.empty()

    excel_results = [
        result
        for result in results
        if result.get("excel_output")
    ]

    for result_index, result in enumerate(results):
        st.divider()
        st.subheader(result["filename"])

        if result.get("error"):
            st.error(f"Não foi possível processar o arquivo: {result['error']}")
            continue

        if not result["pages"]:
            st.warning(
                "Nenhuma página contendo ângulo entre 0 e 360 "
                "associado a NM foi encontrada."
            )
            continue

        human_pages = [page + 1 for page in result["pages"]]
        st.success(
            "Páginas NM encontradas: "
            + ", ".join(str(page) for page in human_pages)
        )

        for page_number, angles in zip(human_pages, result["evidence"]):
            st.write(f"Página {page_number}: {angles}")

        if result["photos"]:
            st.info(f"Fotografias associadas aos ângulos: {len(result['photos'])}")

            preview_labels = [photo["label"] for photo in result["photos"]]
            st.caption("Ângulos organizados: " + " | ".join(preview_labels))

            st.download_button(
                label="Baixar Excel com as fotos organizadas",
                data=result["excel_output"],
                file_name=result["excel_name"],
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                key=f"excel_{result_index}",
            )
        else:
            st.warning(
                "As páginas NM foram localizadas, mas não foi possível associar "
                "as fotografias às legendas."
            )

    if len(excel_results) > 1:
        st.divider()
        zip_buffer = io.BytesIO()

        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            for result in excel_results:
                archive.writestr(result["excel_name"], result["excel_output"])

        st.download_button(
            label="Baixar todos os arquivos Excel em ZIP",
            data=zip_buffer.getvalue(),
            file_name="Fotos_360_Excel.zip",
            mime="application/zip",
            key="all_excel_zip",
        )

st.divider()
st.caption(
    "O Excel é montado dinamicamente com os ângulos efetivamente encontrados. "
    "Os ângulos não precisam seguir intervalos fixos de 30 graus."
)
